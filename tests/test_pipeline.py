import io

import openpyxl
import pytest

from mslog import export, review, service, verify


def open_findings(conn, ftype=None):
    sql = "SELECT * FROM findings WHERE status='open'" + (" AND type=?" if ftype else "")
    return conn.execute(sql, (ftype,) if ftype else ()).fetchall()


def test_detects_planted_errors(loaded):
    qty = open_findings(loaded, "수량불일치")
    assert len(qty) == 1
    assert (qty[0]["base_qty"], qty[0]["compare_qty"], qty[0]["diff"]) == (12000, 11800, 200)

    missing = open_findings(loaded, "자료누락")
    assert [m["vehicle_no"] for m in missing if "서울" in m["vehicle_no"]] and "계근표가 없음" in \
        next(m for m in missing if "서울" in m["vehicle_no"])["detail"]

    assert len(open_findings(loaded, "중복의심")) == 2          # 운행일지 쪽, 계근표 쪽
    nm = open_findings(loaded, "매칭실패")
    assert any("코리아에너지" in f["detail"] for f in nm)

    unc = open_findings(loaded, "인식불확실")
    assert any("불일치" in f["detail"] for f in unc)             # 총중량-공차≠순중량


def test_alias_matching_normalizes_customer_and_item(loaded):
    r = loaded.execute(
        "SELECT * FROM records WHERE doc_type='운행일지' AND customer_raw='대한엘피지'").fetchone()
    assert r["customer_code"] == "C002" and r["item"] == "프로판"
    w = loaded.execute("SELECT * FROM records WHERE doc_type='계근표' AND vehicle_no='경기12가3456' "
                       "AND delivery_date='2026-10-13' AND customer_code='C001'").fetchone()
    assert w is not None                                          # '경기 12가 3456' 공백 제거


def test_duplicate_file_rejected(loaded, samples, upload_dir):
    doc_id, msg = service.register_file(loaded, upload_dir, "복사본.xlsx",
                                        (samples / "운행일지_샘플.xlsx").read_bytes(), "운행일지", "t")
    assert doc_id is None and "이미 등록" in msg


def test_cannot_approve_with_open_error_until_resolved(loaded):
    rid = loaded.execute(
        "SELECT r.id FROM records r JOIN finding_records fr ON fr.record_id=r.id JOIN findings f ON f.id=fr.finding_id "
        "WHERE f.type='수량불일치' AND r.doc_type='계근표'").fetchone()["id"]
    with pytest.raises(review.ReviewError, match="수량불일치"):
        review.approve_record(loaded, rid, "홍길동")
    # 원본 확인 후 계근표 수량을 정정하면 오류가 자동 해소되고 승인 가능
    review.save_record(loaded, rid, {"quantity": "12000"}, "홍길동")
    assert not open_findings(loaded, "수량불일치")
    review.approve_record(loaded, rid, "홍길동")
    edit = loaded.execute("SELECT * FROM edits WHERE record_id=?", (rid,)).fetchone()
    assert (edit["field"], edit["old_value"], edit["new_value"], edit["editor"]) == ("quantity", "11800.0", "12000.0", "홍길동")


def test_resolve_requires_reason(loaded):
    f = open_findings(loaded, "수량불일치")[0]
    with pytest.raises(ValueError):
        verify.resolve_finding(loaded, f["id"], "홍길동", "  ")
    verify.resolve_finding(loaded, f["id"], "홍길동", "원본 확인: 계근표 기준이 맞음")
    assert not open_findings(loaded, "수량불일치")
    verify.run_verification(loaded)                               # 재검증해도 처리 결과 유지
    assert loaded.execute("SELECT status FROM findings WHERE id=?", (f["id"],)).fetchone()[0] == "resolved"


def test_ai_original_preserved_and_edit_unflags_field(loaded):
    rid = loaded.execute("SELECT id FROM records WHERE customer_raw='코리아에너지' AND doc_type='운행일지'").fetchone()["id"]
    review.save_record(loaded, rid, {"customer_raw": "서해에너지"}, "t")
    r = loaded.execute("SELECT * FROM records WHERE id=?", (rid,)).fetchone()
    assert r["customer_code"] == "C003" and "코리아에너지" in r["ai_json"]


def test_export_only_contains_verified(loaded):
    # 아무것도 승인하지 않으면 ERP 입력 대상 없음
    assert export.verified_rows(loaded) == []
    # 서해에너지 10/13 건(운행일지·계근표 일치)만 승인
    ids = [r["id"] for r in loaded.execute(
        "SELECT id FROM records WHERE customer_code='C003' AND delivery_date='2026-10-13'")]
    assert len(ids) == 2
    for i in ids:
        review.approve_record(loaded, i, "홍길동")
    rows = export.verified_rows(loaded)
    assert len(rows) == 1 and rows[0]["quantity"] == 8000 and rows[0]["customer_code"] == "C003"

    wb = openpyxl.load_workbook(io.BytesIO(export.build_workbook(loaded).read()))
    assert wb.sheetnames == ["ERP입력", "일자별현황", "차량별실적", "거래처사업장별실적", "월말누적", "오류·미검증"]
    assert wb["ERP입력"].max_row == 2
    assert wb["오류·미검증"].max_row > 5


def test_unit_conversion_ton_vs_kg(conn):
    from mslog import ingest
    from mslog.db import now
    conn.execute("INSERT INTO documents(doc_type,file_name,stored_name,file_hash,source_kind,uploaded_at,ocr_status)"
                 " VALUES('운행일지','a','a','h1','excel',?,'완료'),('계근표','b','b','h2','excel',?,'완료')",
                 (now(), now()))
    base = {"delivery_date": "2026-10-20", "vehicle_no": "12가3456", "customer": "대한LPG", "item": "프로판"}
    ingest.add_record(conn, 1, "운행일지", {**base, "quantity": 12, "unit": "톤"})
    ingest.add_record(conn, 2, "계근표", {**base, "quantity": 12000, "unit": "kg"})
    conn.commit()
    verify.run_verification(conn)
    assert not open_findings(conn, "수량불일치")
