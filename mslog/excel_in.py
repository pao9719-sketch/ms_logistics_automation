"""엑셀(.xlsx) 업로드 파싱. 헤더 행을 자동으로 찾아 표준 필드로 매핑한다."""
import re

import openpyxl

HEADER_ALIASES = {
    "delivery_date": ["일자", "납품일자", "날짜", "계근일자", "운행일자", "납품일", "계근일"],
    "vehicle_no": ["차량번호", "차량", "차번호", "차번", "차량no"],
    "customer": ["거래처", "거래처명", "납품처", "상호", "업체명"],
    "site": ["사업장", "사업장명", "현장", "납품장소"],
    "item": ["품목", "품목명", "품명", "제품"],
    "quantity": ["순중량", "납품수량", "수량", "실중량", "인수량", "납품량"],
    "unit": ["단위"],
    "gross_weight": ["총중량", "총중", "입차중량"],
    "tare_weight": ["공차중량", "공차", "차량중량", "출차중량"],
}

_LOOKUP = {}
for _field, _names in HEADER_ALIASES.items():
    for _n in _names:
        _LOOKUP[re.sub(r"\s", "", _n).lower()] = _field


def _field_for(header):
    if header is None:
        return None
    return _LOOKUP.get(re.sub(r"\s", "", str(header)).lower())


def read_rows(path):
    """시트의 데이터 행을 [{field: value}]로 반환. 헤더를 찾지 못하면 ValueError."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            for hi, row in enumerate(rows[:15]):
                mapping = {}
                for ci, h in enumerate(row):
                    f = _field_for(h)
                    if f and f not in mapping.values():
                        mapping[ci] = f
                fields = set(mapping.values())
                if len(fields) >= 3 and "quantity" in fields and "delivery_date" in fields:
                    out = []
                    for r in rows[hi + 1:]:
                        rec = {f: r[ci] for ci, f in mapping.items() if ci < len(r)}
                        if all(rec.get(k) in (None, "") for k in ("delivery_date", "vehicle_no", "quantity")):
                            continue
                        out.append(rec)
                    return out
    finally:
        wb.close()
    raise ValueError("헤더 행을 찾지 못했습니다. 일자·차량번호·거래처·품목·수량 열이 필요합니다.")


def read_table(path):
    """기준정보 업로드용: 첫 시트를 [{헤더: 값}]으로."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(c not in (None, "") for c in r)]
