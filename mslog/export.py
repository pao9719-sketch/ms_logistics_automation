"""검증·승인된 데이터의 엑셀 출력.

ERP 입력 대상 = (운행일지·계근표가 모두 있고) 구성 기록이 전부 '검수완료'이며
열린 오류가 없는 납품 건. 수량은 계근표 합계를 사용한다.
그 외는 '오류·미검증' 시트로 분리한다.
"""
import io
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .verify import LOG, WEIGH, build_groups, sum_quantity


def _records(conn, d_from, d_to):
    sql = "SELECT * FROM records WHERE review_status!='제외'"
    args = []
    if d_from:
        sql += " AND delivery_date>=?"
        args.append(d_from)
    if d_to:
        sql += " AND delivery_date<=?"
        args.append(d_to)
    return [dict(r) for r in conn.execute(sql, args)]


def verified_rows(conn, d_from=None, d_to=None):
    recs = _records(conn, d_from, d_to)
    blocked = {r["record_id"] for r in conn.execute(
        "SELECT fr.record_id FROM finding_records fr JOIN findings f ON f.id=fr.finding_id WHERE f.status='open'")}
    rows = []
    for key, g in build_groups(recs).items():
        logs, weighs = g[LOG], g[WEIGH]
        members = logs + weighs
        if not logs or not weighs:
            continue
        if any(m["review_status"] != "검수완료" or m["id"] in blocked for m in members):
            continue
        qty, unit = sum_quantity(weighs)
        first = weighs[0]
        rows.append({
            "date": key[0], "vehicle": first["vehicle_no"],
            "customer_code": first["customer_code"] or "", "customer_name": first["customer_name"] or first["customer_raw"],
            "site": next((m["site"] for m in weighs + logs if m["site"]), ""),
            "item": first["item"], "quantity": qty, "unit": unit,
            "n_log": len(logs), "n_weigh": len(weighs),
        })
    rows.sort(key=lambda r: (r["date"], r["vehicle"] or "", r["customer_name"] or "", r["item"] or ""))
    return rows


def unverified_rows(conn, d_from=None, d_to=None):
    """열린 오류 목록 + 아직 검수되지 않은 기록(오류 없음 포함)."""
    out = []
    for f in conn.execute("SELECT * FROM findings WHERE status='open' ORDER BY delivery_date, vehicle_no"):
        if d_from and (f["delivery_date"] or "") < d_from:
            continue
        if d_to and (f["delivery_date"] or "9999") > d_to:
            continue
        out.append(("오류", f["type"], f["delivery_date"], f["vehicle_no"], f["customer"], f["item"],
                    f["base_qty"], f["compare_qty"], f["detail"], f["assignee"] or ""))
    for r in _records(conn, d_from, d_to):
        if r["review_status"] == "미검수":
            out.append(("미검수", "검수 대기", r["delivery_date"], r["vehicle_no"],
                        r["customer_name"] or r["customer_raw"], r["item"], r["quantity"], None,
                        f"{r['doc_type']} 문서 #{r['document_id']}", ""))
    return out


def _agg(rows, keyfunc):
    acc = defaultdict(lambda: [0.0, 0])
    for r in rows:
        k = keyfunc(r)
        acc[k][0] += r["quantity"]
        acc[k][1] += 1
    return acc


def _sheet(wb, title, headers, rows, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
        c.alignment = Alignment(horizontal="center")
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if r[i - 1] is not None])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 4, 10), 50)
    ws.freeze_panes = "A2"
    return ws


def build_workbook(conn, d_from=None, d_to=None):
    rows = verified_rows(conn, d_from, d_to)
    wb = Workbook()
    _sheet(wb, "ERP입력", [h for h, _ in config.ERP_COLUMNS],
           [[r[k] for _, k in config.ERP_COLUMNS] for r in rows], first=True)

    daily = _agg(rows, lambda r: (r["date"], r["item"], r["unit"]))
    _sheet(wb, "일자별현황", ["납품일자", "품목", "단위", "건수", "수량합계"],
           [[k[0], k[1], k[2], v[1], v[0]] for k, v in sorted(daily.items(), key=lambda x: tuple(map(str, x[0])))])

    veh = _agg(rows, lambda r: (r["vehicle"], r["item"], r["unit"]))
    _sheet(wb, "차량별실적", ["차량번호", "품목", "단위", "건수", "수량합계"],
           [[k[0], k[1], k[2], v[1], v[0]] for k, v in sorted(veh.items(), key=lambda x: tuple(map(str, x[0])))])

    cust = _agg(rows, lambda r: (r["customer_code"], r["customer_name"], r["site"], r["item"], r["unit"]))
    _sheet(wb, "거래처사업장별실적", ["거래처코드", "거래처명", "사업장", "품목", "단위", "건수", "수량합계"],
           [[*k, v[1], v[0]] for k, v in sorted(cust.items(), key=lambda x: tuple(map(str, x[0])))])

    month = _agg(rows, lambda r: (r["date"][:7], r["customer_code"], r["customer_name"], r["site"], r["item"], r["unit"]))
    _sheet(wb, "월말누적", ["월", "거래처코드", "거래처명", "사업장", "품목", "단위", "건수", "수량합계"],
           [[*k, v[1], v[0]] for k, v in sorted(month.items(), key=lambda x: tuple(map(str, x[0])))])

    _sheet(wb, "오류·미검증",
           ["구분", "유형", "납품일자", "차량번호", "거래처", "품목", "수량(기준)", "수량(비교)", "내용", "담당자"],
           [list(r) for r in unverified_rows(conn, d_from, d_to)])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
